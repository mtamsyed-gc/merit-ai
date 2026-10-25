"""Gate engine semantics: outcome rules, status validation, sign-offs, crash safety, report files."""
import json
import tempfile
from pathlib import Path

from helpers import done, ok, tmpdir
from openpyxl import load_workbook

from meritai import gates
from meritai.gates import Result

gates.built()                       # register real checks first
saved = dict(gates.REGISTRY)
G7 = ["G7.1", "G7.2", "G7.3", "G7.4", "G7.5"]
G1 = ["G1.1", "G1.2", "G1.3", "G1.4", "G1.5", "G1.6", "G1.7", "G1.8"]


def with_impl(impls):
    gates.REGISTRY.clear()
    gates.REGISTRY.update(saved)
    for cid, st in impls.items():
        if st == "CRASH":
            def boom(ctx):
                raise ValueError("Original Line: C01|NAS|P001|M001,784-1990-1234567-1")
            gates.REGISTRY[cid] = boom
        else:
            gates.REGISTRY[cid] = (lambda s, c: (lambda ctx: Result(c, s, f"{s} test", value=1)))(st, cid)


with_impl({c: "PASS" for c in G7})
ok(gates.run_gate("G7", {}, "s").outcome == "PASS", "all pass -> PASS")
with_impl({**{c: "PASS" for c in G7}, "G7.2": "STOP"})
ok(gates.run_gate("G7", {}, "s").outcome == "STOP", "one Stop -> STOP")
with_impl({c: "PASS" for c in G7[:4]})
ok(gates.run_gate("G7", {}, "s").outcome == "INCOMPLETE", "a check not built -> INCOMPLETE")
ok(gates.run_gate("G7", {}, "s", allow_incomplete=True).outcome == "PASS (incomplete)",
   "a dev run with unbuilt checks is labelled incomplete, never plain PASS")
with_impl({**{c: "PASS" for c in G7}, "G7.1": "CRASH"})
r = gates.run_gate("G7", {}, "s")
ok(r.outcome == "STOP" and r.results[0].status == "STOP", "a crashing check is a Stop, never a pass")
ok("784-1990" not in r.results[0].message and "Original Line" not in r.results[0].message,
   "a crash message never carries the exception text (raw data)")
with_impl({**{c: "PASS" for c in G7}, "G7.3": "INFO"})
ok(gates.run_gate("G7", {}, "s").outcome == "PASS", "Info never blocks")
for bad in ("Pass", "STOP ", "OK", None):
    with_impl({**{c: "PASS" for c in G7}, "G7.4": bad})
    rr = gates.run_gate("G7", {}, "s")
    ok(rr.outcome == "STOP", f"unknown status {bad!r} -> STOP")

# Flags and sign-offs (G1.7 is a Flag check, signed by the Pool owner)
with_impl({**{c: "PASS" for c in G1}, "G1.7": "FLAG"})
r = gates.run_gate("G1", {}, "C99")
ok(r.outcome == "FLAG", "Flag without sign-off -> FLAG")
fp = [x for x in r.results if x.check_id == "G1.7"][0].fingerprint()
so = tmpdir('t_') / "signoffs.json"
try:
    gates.add_signoff(so, "G1.7", "C99", fp, "A. Person", "Lead actuary", "wrong role")
    ok(False, "sign-off in the wrong role refused")
except ValueError:
    ok(True, "sign-off in the wrong role refused")
for args in (("", "reason"), ("A. Person", "")):
    try:
        gates.add_signoff(so, "G1.7", "C99", fp, args[0], "Pool owner", args[1])
        ok(False, f"sign-off needs signer and reason {args}")
    except ValueError:
        ok(True, f"sign-off needs signer and reason {args}")
gates.add_signoff(so, "G1.7", "C99", fp, "A. Person", "Pool owner", "replacement agreed with the cedant")
s = gates.load_signoffs(so)
r = gates.run_gate("G1", {}, "C99", signoffs=s)
ok(r.outcome == "PASS", "signed Flag -> PASS")
ok(gates.run_gate("G1", {}, "OTHER", signoffs=s).outcome == "FLAG", "a sign-off covers its own subject only")
with_impl({**{c: "PASS" for c in G1}, "G1.7": "FLAG"})
gates.REGISTRY["G1.7"] = lambda ctx: Result("G1.7", "FLAG", "a different finding", value=2)
ok(gates.run_gate("G1", {}, "C99", signoffs=s).outcome == "FLAG", "a changed finding needs a new sign-off")
with_impl({**{c: "PASS" for c in G1}, "G1.7": "FLAG"})
hand = [{"check_id": "G1.7", "subject": "C99", "fingerprint": fp, "signer": "", "role": "Pool owner", "reason": "",
         "date": "2026-10-25"}, {"check_id": "G1.7"}, "not a row"]
rr = gates.run_gate("G1", {}, "C99", signoffs=hand)
ok(rr.outcome == "FLAG" and rr.ignored_signoffs == 1, "hand-edited blank or malformed sign-offs are ignored and counted")
with_impl({**{c: "PASS" for c in G1}, "G1.1": "STOP"})
fp_stop = Result("G1.1", "STOP", "STOP test", value=1, subject="C99").fingerprint()
stop_so = [{"check_id": "G1.1", "subject": "C99", "fingerprint": fp_stop, "signer": "x", "role": "Pool owner",
            "reason": "y", "date": "2026-10-25"}]
ok(gates.run_gate("G1", {}, "C99", signoffs=stop_so).outcome == "STOP", "a STOP is never signed away")

# report
with_impl({**{c: "PASS" for c in G1}, "G1.7": "FLAG"})
r = gates.run_gate("G1", {}, "C99", signoffs=s)
out = tmpdir('t_')
xl, md = gates.write_report([r], out, "t")
wb = load_workbook(xl)
ok(wb.sheetnames == ["Summary", "Checks", "Details"], "report workbook has Summary, Checks, Details")
rows = list(wb["Checks"].iter_rows(values_only=True))
ok(rows[0][:4] == ("Subject", "Check", "Title", "Status"), "Checks sheet header")
ok(any(x[3] == "FLAG (signed)" for x in rows[1:]), "signed Flag shown as such")
hdr = rows[0]
ok(all(x[hdr.index("What it means")] for x in rows[1:]), "every row carries the plain-language meaning")
ok(any(x[hdr.index("Fingerprint")] == fp for x in rows[1:]), "Flag rows show the fingerprint to sign")
ok("G1 Ingest" in md.read_text(encoding="utf-8"), "markdown report written")
xl2, _ = gates.write_report([r], tmpdir('t_'), "t")
ok(xl.read_bytes() == xl2.read_bytes(), "same results -> byte-identical report workbook")
gates.REGISTRY.clear()
gates.REGISTRY.update(saved)
done("test_gates")
